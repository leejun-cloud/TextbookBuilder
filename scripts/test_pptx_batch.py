#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
test_pptx_batch.py — pptx_batch.py 자체 테스트 (외부 프레임워크 불필요)

실행:
    python3 scripts/test_pptx_batch.py
"""

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

SCRIPT = Path(__file__).resolve().parent / "pptx_batch.py"


def make_deck(path, small_pt=12, big_pt=24):
    prs = Presentation()
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)

    # 12pt / 24pt 텍스트박스
    tb1 = slide.shapes.add_textbox(Cm(1), Cm(1), Cm(10), Cm(2))
    p1 = tb1.text_frame.paragraphs[0]
    r1 = p1.add_run()
    r1.text = "작은 글씨"
    r1.font.size = Pt(small_pt)

    tb2 = slide.shapes.add_textbox(Cm(1), Cm(4), Cm(10), Cm(2))
    p2 = tb2.text_frame.paragraphs[0]
    r2 = p2.add_run()
    r2.text = "큰 글씨"
    r2.font.size = Pt(big_pt)

    # run 경계를 넘나드는 문구: "삭제할 " + "문구입니다" -> "삭제할 문구"가 두 run에 걸침
    tb3 = slide.shapes.add_textbox(Cm(1), Cm(7), Cm(14), Cm(2))
    p3 = tb3.text_frame.paragraphs[0]
    ra = p3.add_run()
    ra.text = "삭제할 "
    ra.font.size = Pt(14)
    rb = p3.add_run()
    rb.text = "문구입니다 끝"
    rb.font.size = Pt(14)

    # replace-map 테스트용: 한 문단에 여러 개의 독립된 찾을 문구
    tb4 = slide.shapes.add_textbox(Cm(1), Cm(13), Cm(16), Cm(2))
    p4 = tb4.text_frame.paragraphs[0]
    r4 = p4.add_run()
    r4.text = "핵심질문: 하나님은 누구신가? 오늘의목표: 예배자로 살기"
    r4.font.size = Pt(14)

    # replace-map 부분문자열 순서 테스트용: "질문" 이 "질문입니다"의 substring
    tb5 = slide.shapes.add_textbox(Cm(1), Cm(16), Cm(16), Cm(2))
    p5 = tb5.text_frame.paragraphs[0]
    r5 = p5.add_run()
    r5.text = "질문입니다 이렇게 답하세요"
    r5.font.size = Pt(14)

    # 표
    rows, cols = 2, 2
    tbl_shape = slide.shapes.add_table(rows, cols, Cm(1), Cm(10), Cm(10), Cm(3))
    table = tbl_shape.table
    table.cell(0, 0).text = "표머리1"
    table.cell(0, 1).text = "표머리2"
    table.cell(1, 0).text = "표내용1"
    table.cell(1, 1).text = "표내용2"

    # FF0000 채우기 도형
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Cm(16), Cm(1), Cm(4), Cm(3))
    sp.fill.solid()
    sp.fill.fore_color.rgb = RGBColor.from_string("FF0000")

    prs.save(str(path))


def run_cli(*args, cwd=None):
    cmd = [sys.executable, str(SCRIPT)] + list(args)
    proc = subprocess.run(cmd, capture_output=True, text=True, cwd=cwd)
    return proc


def assert_true(cond, msg):
    if not cond:
        raise AssertionError(msg)


def test_list(tmp, deck_a, deck_b):
    proc = run_cli("list", str(deck_a), str(deck_b))
    assert_true(proc.returncode == 0, f"list 실패: {proc.stdout}\n{proc.stderr}")
    assert_true("전체 텍스트 런" in proc.stdout, "list 출력에 인벤토리 정보 없음")
    print("[PASS] list")


def test_scale_font(tmp, deck_a):
    out_dir = tmp / "out_scale"
    proc = run_cli("scale-font", "--factor", "2", str(deck_a), "--out-dir", str(out_dir))
    assert_true(proc.returncode == 0, f"scale-font 실패: {proc.stdout}\n{proc.stderr}")
    out_path = out_dir / deck_a.name
    assert_true(out_path.exists(), "scale-font 출력 파일 없음")
    prs = Presentation(str(out_path))
    slide = prs.slides[0]
    sizes = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size is not None:
                        sizes.append(run.font.size.pt)
    assert_true(24 in sizes, f"12pt -> 24pt 배율 조정 실패: {sizes}")
    assert_true(48 in sizes, f"24pt -> 48pt 배율 조정 실패: {sizes}")
    print("[PASS] scale-font")


def test_min_font(tmp, deck_a):
    out_dir = tmp / "out_min"
    proc = run_cli("min-font", "--size", "16", str(deck_a), "--out-dir", str(out_dir))
    assert_true(proc.returncode == 0, f"min-font 실패: {proc.stdout}\n{proc.stderr}")
    out_path = out_dir / deck_a.name
    prs = Presentation(str(out_path))
    slide = prs.slides[0]
    sizes = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size is not None:
                        sizes.append(run.font.size.pt)
    assert_true(16 in sizes, f"12pt -> 16pt 상향 실패: {sizes}")
    assert_true(24 in sizes, f"24pt(변경 불필요)이 사라짐: {sizes}")
    assert_true(12 not in sizes, f"12pt 가 그대로 남음: {sizes}")
    print("[PASS] min-font")


def _get_split_phrase_shape(slide):
    for shape in slide.shapes:
        if shape.has_text_frame and "문구" in shape.text_frame.text:
            return shape
    return None


def test_replace_text_spanning_runs(tmp, deck_a):
    out_dir = tmp / "out_replace"
    proc = run_cli("replace-text", "--find", "삭제할 문구", "--replace", "새 문구",
                   str(deck_a), "--out-dir", str(out_dir))
    assert_true(proc.returncode == 0, f"replace-text 실패: {proc.stdout}\n{proc.stderr}")
    out_path = out_dir / deck_a.name
    prs = Presentation(str(out_path))
    shape = _get_split_phrase_shape(prs.slides[0])
    assert_true(shape is not None, "치환 대상 도형을 찾지 못함")
    joined = shape.text_frame.text
    assert_true("삭제할 문구" not in joined, f"원본 문구가 여전히 존재: {joined!r}")
    assert_true("새 문구" in joined, f"치환된 문구가 없음: {joined!r}")
    print("[PASS] replace-text (run 경계를 넘는 치환)")


def test_remove_text(tmp, deck_a):
    out_dir = tmp / "out_remove"
    proc = run_cli("remove-text", "--find", "삭제할 문구",
                   str(deck_a), "--out-dir", str(out_dir))
    assert_true(proc.returncode == 0, f"remove-text 실패: {proc.stdout}\n{proc.stderr}")
    out_path = out_dir / deck_a.name
    prs = Presentation(str(out_path))
    shape = None
    for s in prs.slides[0].shapes:
        if s.has_text_frame and "입니다" in s.text_frame.text:
            shape = s
            break
    assert_true(shape is not None, "대상 도형을 찾지 못함")
    joined = shape.text_frame.text
    assert_true("삭제할 문구" not in joined, f"문구 제거 실패: {joined!r}")
    print("[PASS] remove-text")


def _write_pairs(tmp, name, pairs):
    path = tmp / name
    path.write_text(json.dumps(pairs, ensure_ascii=False), encoding="utf-8")
    return path


def _find_shape_with_text(slide, needle):
    for shape in slide.shapes:
        if shape.has_text_frame and needle in shape.text_frame.text:
            return shape
    return None


def test_replace_map_multi_pair(tmp, deck_a):
    pairs_path = _write_pairs(tmp, "pairs_multi.json", [
        {"find": "핵심질문", "replace": "주질문"},
        {"find": "오늘의목표", "replace": "오늘목표"},
    ])
    out_dir = tmp / "out_map_multi"
    proc = run_cli("replace-map", "--map", str(pairs_path), str(deck_a), "--out-dir", str(out_dir))
    assert_true(proc.returncode == 0, f"replace-map 실패: {proc.stdout}\n{proc.stderr}")
    out_path = out_dir / deck_a.name
    prs = Presentation(str(out_path))
    shape = _find_shape_with_text(prs.slides[0], "주질문")
    assert_true(shape is not None, "첫 번째 쌍이 반영되지 않음")
    joined = shape.text_frame.text
    assert_true("핵심질문" not in joined, f"첫 번째 쌍의 원본 문구가 남음: {joined!r}")
    assert_true("오늘목표" in joined, f"두 번째 쌍이 반영되지 않음: {joined!r}")
    assert_true("오늘의목표" not in joined, f"두 번째 쌍의 원본 문구가 남음: {joined!r}")
    print("[PASS] replace-map (한 파일에 여러 쌍을 한 번에 적용)")


def test_replace_map_zero_hit_reported(tmp, deck_a):
    pairs_path = _write_pairs(tmp, "pairs_zero.json", [
        {"find": "핵심질문", "replace": "주질문"},
        {"find": "존재하지않는문구", "replace": "X"},
    ])
    out_dir = tmp / "out_map_zero"
    proc = run_cli("replace-map", "--map", str(pairs_path), str(deck_a), "--out-dir", str(out_dir))
    assert_true(proc.returncode == 0, f"replace-map 실패: {proc.stdout}\n{proc.stderr}")
    assert_true("미발견" in proc.stdout, f"0건 쌍이 미발견으로 보고되지 않음:\n{proc.stdout}")
    assert_true("존재하지않는문구" in proc.stdout, f"미발견 문구 자체가 보고되지 않음:\n{proc.stdout}")
    print("[PASS] replace-map (0건 쌍은 미발견으로 보고)")


def test_replace_map_longer_find_wins(tmp, deck_a):
    pairs_path = _write_pairs(tmp, "pairs_substr.json", [
        {"find": "질문", "replace": "Q"},
        {"find": "질문입니다", "replace": "QQQ"},
    ])
    out_dir = tmp / "out_map_substr"
    proc = run_cli("replace-map", "--map", str(pairs_path), str(deck_a), "--out-dir", str(out_dir))
    assert_true(proc.returncode == 0, f"replace-map 실패: {proc.stdout}\n{proc.stderr}")
    out_path = out_dir / deck_a.name
    prs = Presentation(str(out_path))
    shape = _find_shape_with_text(prs.slides[0], "이렇게 답하세요")
    assert_true(shape is not None, "대상 도형을 찾지 못함")
    joined = shape.text_frame.text
    assert_true("QQQ" in joined, f"긴 문구('질문입니다')가 우선 적용되지 않음: {joined!r}")
    assert_true("질문" not in joined, f"짧은 문구가 남거나 잘못 치환됨: {joined!r}")
    print("[PASS] replace-map (긴 find 문구가 짧은 부분문자열보다 먼저 적용됨)")


def test_replace_color(tmp, deck_a):
    out_dir = tmp / "out_color"
    proc = run_cli("replace-color", "--from", "FF0000", "--to", "0000FF",
                   str(deck_a), "--out-dir", str(out_dir))
    assert_true(proc.returncode == 0, f"replace-color 실패: {proc.stdout}\n{proc.stderr}")
    out_path = out_dir / deck_a.name
    prs = Presentation(str(out_path))
    found_blue = False
    for shape in prs.slides[0].shapes:
        if shape.shape_type is not None and getattr(shape, "fill", None) is not None:
            try:
                if shape.fill.type is not None and str(shape.fill.fore_color.rgb) == "0000FF":
                    found_blue = True
            except Exception:
                pass
    assert_true(found_blue, "FF0000 -> 0000FF 색상 치환 실패")
    print("[PASS] replace-color")


def test_dry_run(tmp, deck_a):
    out_dir = tmp / "out_dry"
    proc = run_cli("scale-font", "--factor", "3", str(deck_a), "--out-dir", str(out_dir), "--dry-run")
    assert_true(proc.returncode == 0, f"dry-run 실패: {proc.stdout}\n{proc.stderr}")
    assert_true(not out_dir.exists() or not any(out_dir.iterdir()), "dry-run 인데 출력 파일이 생성됨")
    print("[PASS] --dry-run (아무 것도 쓰지 않음)")


def test_in_place_backup(tmp, deck_a):
    work_copy = tmp / "inplace_test.pptx"
    shutil.copy2(deck_a, work_copy)
    proc = run_cli("min-font", "--size", "20", str(work_copy), "--in-place")
    assert_true(proc.returncode == 0, f"--in-place 실패: {proc.stdout}\n{proc.stderr}")

    backup_dirs = [d for d in work_copy.parent.iterdir() if d.is_dir() and d.name.startswith("백업_")]
    assert_true(len(backup_dirs) == 1, f"백업 폴더가 생성되지 않음: {list(work_copy.parent.iterdir())}")
    backup_file = backup_dirs[0] / work_copy.name
    assert_true(backup_file.exists(), "백업 폴더 안에 원본 파일이 없음")

    # 백업본은 수정 전 원본(12pt 포함)이어야 함
    prs_backup = Presentation(str(backup_file))
    backup_sizes = []
    for shape in prs_backup.slides[0].shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size is not None:
                        backup_sizes.append(run.font.size.pt)
    assert_true(12 in backup_sizes, f"백업본이 원본이 아님: {backup_sizes}")

    # 실제 파일은 수정된 버전(20pt 포함, 12pt 없음)이어야 함
    prs_modified = Presentation(str(work_copy))
    modified_sizes = []
    for shape in prs_modified.slides[0].shapes:
        if shape.has_text_frame:
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size is not None:
                        modified_sizes.append(run.font.size.pt)
    assert_true(12 not in modified_sizes, f"--in-place 가 원본을 수정하지 않음: {modified_sizes}")
    print("[PASS] --in-place (자동 백업 생성)")


def test_extract(tmp, deck_a):
    out_dir = tmp / "out_extract"
    proc = run_cli("extract", "--format", "md", str(deck_a), "--out-dir", str(out_dir))
    assert_true(proc.returncode == 0, f"extract 실패: {proc.stdout}\n{proc.stderr}")
    md_path = out_dir / f"{deck_a.stem}.md"
    assert_true(md_path.exists(), "extract 결과 md 파일 없음")
    text = md_path.read_text(encoding="utf-8")
    assert_true("슬라이드 1" in text, "슬라이드 번호 표시 없음")
    assert_true("작은 글씨" in text, "본문 텍스트 누락")
    assert_true("표머리1" in text, "표 내용 누락")
    print("[PASS] extract")


def test_missing_output_flag_errors(tmp, deck_a):
    proc = run_cli("min-font", "--size", "10", str(deck_a))
    assert_true(proc.returncode != 0, "출력 방식 없이 실행했는데 오류가 발생하지 않음")
    print("[PASS] 출력 플래그 누락 시 오류")


def main():
    tmp_root = Path(tempfile.mkdtemp(prefix="pptx_batch_test_"))
    try:
        deck_a = tmp_root / "deckA.pptx"
        deck_b = tmp_root / "deckB.pptx"
        make_deck(deck_a)
        make_deck(deck_b, small_pt=11, big_pt=22)

        tests = [
            lambda: test_list(tmp_root, deck_a, deck_b),
            lambda: test_scale_font(tmp_root, deck_a),
            lambda: test_min_font(tmp_root, deck_a),
            lambda: test_replace_text_spanning_runs(tmp_root, deck_a),
            lambda: test_remove_text(tmp_root, deck_a),
            lambda: test_replace_map_multi_pair(tmp_root, deck_a),
            lambda: test_replace_map_zero_hit_reported(tmp_root, deck_a),
            lambda: test_replace_map_longer_find_wins(tmp_root, deck_a),
            lambda: test_replace_color(tmp_root, deck_a),
            lambda: test_dry_run(tmp_root, deck_a),
            lambda: test_in_place_backup(tmp_root, deck_a),
            lambda: test_extract(tmp_root, deck_a),
            lambda: test_missing_output_flag_errors(tmp_root, deck_a),
        ]

        failed = 0
        for t in tests:
            try:
                t()
            except AssertionError as e:
                failed += 1
                print(f"[FAIL] {e}")
            except Exception as e:
                failed += 1
                print(f"[ERROR] {type(e).__name__}: {e}")

        print()
        if failed:
            print(f"=== {failed}건 실패 ===")
            sys.exit(1)
        else:
            print("=== 전체 테스트 통과 ===")
            sys.exit(0)
    finally:
        shutil.rmtree(tmp_root, ignore_errors=True)


if __name__ == "__main__":
    main()
