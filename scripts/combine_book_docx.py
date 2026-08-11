#!/usr/bin/env python3
"""여러 과의 draft.md를 한 권(합본) Word(docx)로 합친다.

catechism_lesson_docx.py의 렌더러를 그대로 재사용하되, 과 경계마다
- 배너 색 순번(position)을 0으로 리셋하고
- 강제 페이지 나누기(page break)를 삽입한다.

사용법: python3 combine_book_docx.py <lessons_dir> -o <out.docx> --start 1 --end 47
  lessons_dir 예: .../lessons/book1  (그 아래 01, 02, ... 폴더에 draft.md가 있다고 가정)
"""
import argparse
import os

import docx
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

import catechism_lesson_docx as base

FONT = base.FONT


def add_page_break(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    br = OxmlElement("w:br")
    br.set(qn("w:type"), "page")
    run._r.append(br)


def render_lesson(doc, md_path):
    md_text = open(md_path, encoding="utf-8").read()
    blocks = base.parse_blocks(md_text)

    position = 0
    for header, body in blocks:
        if header == "COVER":
            base.render_cover(doc, base.parse_kv(body))
        elif header == "오늘의 말씀":
            base.render_scripture(doc, body)
        elif header == "FOOTER":
            base.render_footer_line(doc, body)
        else:
            m = base.SECTION_NUM_RE.match(header)
            title_text = header
            if m:
                position += 1
                title_text = f"{m.group(1)}.  {m.group(2)}"
                topic = m.group(2)
            else:
                topic = header
            base.render_banner(doc, position, title_text)
            if "이 시간의 기대" in topic:
                base.render_expectation(doc, body)
            elif "오프닝" in topic:
                base.render_opening(doc, body)
            elif "깊이 파기" in topic:
                base.render_catechism(doc, body)
            elif "내 삶에 담아가기" in topic:
                base.render_note_and_numbered(doc, body)
            elif "챌린지" in topic:
                base.render_note_and_checkbox(doc, body)
            elif "명언으로 마무리" in topic:
                base.render_quotes(doc, body)
            elif "내 말로 정리하기" in topic:
                base.render_writing_box(doc, body)
            else:
                base.render_note_and_numbered(doc, body)


def build_combined(lessons_dir, out_path, start, end, title=None):
    doc = docx.Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), FONT)

    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    if title:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(200)
        r = p.add_run(title)
        r.font.size = Pt(28)
        r.font.bold = True
        r.font.name = FONT
        rpr2 = r._element.get_or_add_rPr()
        rf2 = OxmlElement("w:rFonts")
        rf2.set(qn("w:eastAsia"), FONT)
        rpr2.append(rf2)
        add_page_break(doc)

    missing = []
    first = True
    for n in range(start, end + 1):
        lesson_dir = os.path.join(lessons_dir, f"{n:02d}")
        md_path = os.path.join(lesson_dir, "draft.md")
        if not os.path.isfile(md_path):
            missing.append(n)
            continue
        if not first:
            add_page_break(doc)
        first = False
        render_lesson(doc, md_path)

    doc.save(out_path)
    print(f"저장됨: {out_path}")
    if missing:
        print(f"경고 — 누락된 과: {missing}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("lessons_dir")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--start", type=int, default=1)
    ap.add_argument("--end", type=int, default=47)
    ap.add_argument("--title", default=None)
    a = ap.parse_args()
    build_combined(a.lessons_dir, a.out, a.start, a.end, a.title)
