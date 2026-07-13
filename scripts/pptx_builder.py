#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
pptx_builder.py — TextbookBuilder 프리미엄 PPT 생성기 (스타일 라이브러리 기반)

사용법:
    python3 scripts/pptx_builder.py spec.json -o output.pptx [--style orange-gradient]

스타일은 ppt-styles/<이름>.json 에서 로드된다 (색·폰트·그림자·모티프 파라미터).
레이아웃 구도(geometry)는 코드에, 색/폰트/질감은 스타일 프로필에 분리되어 있다.

SPEC 형식:
{
  "title": "덱 제목(푸터에 표시)",
  "font": "Pretendard",            # 선택. 미지정 시 스타일의 font.main
  "style": "orange-gradient",      # 선택. CLI --style 이 우선
  "slides": [ { "layout": "<이름>", ...필드..., "notes": "발표자 노트(선택)" } ]
}

레이아웃별 필드 (모든 필드는 문자열, 생략 가능. image 는 파일 경로 — 없으면 그라데이션
플레이스홀더가 그려진다):

  cover       : org, kicker, title, tagline, subtitle, meta
  toc         : kicker, title, image, items[<=6 str]
  intro       : kicker, name, role, text, image
  overview    : kicker, title, blocks[<=3 {head, desc}], cards[<=3 str]
  three_pillars: kicker, title, pillars[<=3 {head, desc, image}]
  team        : kicker, title, members[<=4 {name, role, image}]
  two_column  : kicker, title, columns[<=2 {head, desc}]
  process     : kicker, title, steps[<=3 {head, desc}]
  four_cards  : kicker, title, cards[<=4 {head, desc}]
  vision      : kicker, title, badge, text, image
  gallery     : kicker, title, images[<=3 str 경로], caption
  kpi         : kicker, title, analysis{head, bullets[<=3 str]}, rings[<=3 {value(0-100), label}]
  report      : kicker, title, image, findings[<=4 str]
  bar_chart   : kicker, title, bars[<=6 {label, value(number)}], summary{stat, text}
  full_image  : title, text, image
  testimonial : kicker, title, quotes[<=2 {text, name, role, image}]
  case_study  : kicker, title, cases[<=3 {caption, image}]
  section     : number, title, subtitle
  thank_you   : title, subtitle, contact
"""

import argparse
import copy
import json
import math
import sys
from pathlib import Path

from pptx import Presentation
from pptx.util import Emu, Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

# ---------------------------------------------------------------- 상수/기본 스타일

SLIDE_W_EMU = 12192000
SLIDE_H_EMU = 6858000
SW = Emu(SLIDE_W_EMU).cm   # 33.867 cm
SH = Emu(SLIDE_H_EMU).cm   # 19.05 cm

STYLES_DIR = Path(__file__).resolve().parent.parent / "ppt-styles"

# 스파스한 스타일 프로필이 와도 동작하도록 하는 완전한 기본값
DEFAULT_STYLE = {
    "name": "default",
    "palette": {
        "primary": "FF6A00", "bright": "FF8A00", "accent": "FFC400",
        "charcoal": "222222", "gray": "777777", "light": "F3F3F3",
        "white": "FFFFFF", "divider": "E4E4E4", "footer": "AAAAAA",
        "dark_bg": "1C1C1C", "dark_sub": "D5D5D5",
    },
    "gradient": {
        "main": {"colors": ["FF6A00", "FFC400"], "angle": 0},
        "diagonal": {"colors": ["FF6A00", "FFC400"], "angle": 45},
        "vertical": {"colors": ["FF8A00", "FFC400"], "angle": 90},
        "soft": {"colors": ["FFF3E3", "FFE3B8"], "angle": 0},
        "dark": {"colors": ["1C1C1C", "3A2A18"], "angle": 60},
    },
    "font": {"main": "Pretendard", "fallback": "Malgun Gothic", "latin": "Arial"},
    "type_scale": {
        "cover_title": 54, "title": 30, "section": 28, "subtitle": 14,
        "body": 16, "card_head": 18, "card_body": 12, "caption": 11,
        "kicker": 12, "big_number": 26, "footer": 8.5,
    },
    "margins": {"left_cm": 1.8, "top_cm": 1.25, "right_cm": 1.8, "bottom_cm": 1.1},
    "corner_radius": {"card": 0.038, "bar": 0.1667},
    "card_line": {"color": "E4E4E4", "width_emu": 12700},
    "shadow": {"blur": 177800, "dist": 50800, "dir": 5400000, "alpha": 30000},
    "motif": {
        "ring_thickness": 0.14, "cover_ring_cm": 17.0,
        "wave_height_cm": 2.8, "accent_circles_cm": [1.6, 0.8],
    },
}


def deep_merge(base, override):
    out = copy.deepcopy(base)
    for k, v in (override or {}).items():
        if isinstance(v, dict) and isinstance(out.get(k), dict):
            out[k] = deep_merge(out[k], v)
        else:
            out[k] = v
    return out


def load_style(name):
    path = STYLES_DIR / f"{name}.json"
    if not path.exists():
        avail = sorted(p.stem for p in STYLES_DIR.glob("*.json"))
        sys.exit(f"[오류] 스타일 '{name}'을 찾을 수 없습니다. 사용 가능: {', '.join(avail) or '(없음)'}")
    with open(path, encoding="utf-8") as f:
        return deep_merge(DEFAULT_STYLE, json.load(f))


# ---------------------------------------------------------------- 로우레벨 헬퍼

A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def _a(tag):
    return f"{{{A_NS}}}{tag}"


def col(S, key):
    return RGBColor.from_string(S["palette"][key])


def _spPr(shape):
    return shape._element.spPr


def _clear_fill_xml(spPr):
    for tag in ("noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill"):
        for el in spPr.findall(_a(tag)):
            spPr.remove(el)


def _insert_fill(spPr, fill_el):
    """spPr 내 올바른 위치(geom 뒤, ln 앞)에 fill 요소 삽입"""
    _clear_fill_xml(spPr)
    ln = spPr.find(_a("ln"))
    if ln is not None:
        ln.addprevious(fill_el)
    else:
        eff = spPr.find(_a("effectLst"))
        if eff is not None:
            eff.addprevious(fill_el)
        else:
            spPr.append(fill_el)


def gradient_fill(shape, colors, angle=0, alphas=None):
    """colors: hex 리스트, angle: 도(0=좌→우), alphas: 0-100000 리스트(선택)"""
    n = len(colors)
    gs_items = []
    for i, c in enumerate(colors):
        pos = int(i * 100000 / (n - 1)) if n > 1 else 0
        alpha = f'<a:alpha val="{alphas[i]}"/>' if alphas else ""
        gs_items.append(f'<a:gs pos="{pos}"><a:srgbClr val="{c}">{alpha}</a:srgbClr></a:gs>')
    xml = (f'<a:gradFill xmlns:a="{A_NS}" rotWithShape="1"><a:gsLst>{"".join(gs_items)}</a:gsLst>'
           f'<a:lin ang="{int(angle * 60000)}" scaled="1"/></a:gradFill>')
    _insert_fill(_spPr(shape), etree.fromstring(xml))


def solid_fill_alpha(shape, hexcol, alpha):
    """alpha: 0-100000 (100000=불투명)"""
    xml = (f'<a:solidFill xmlns:a="{A_NS}"><a:srgbClr val="{hexcol}">'
           f'<a:alpha val="{alpha}"/></a:srgbClr></a:solidFill>')
    _insert_fill(_spPr(shape), etree.fromstring(xml))


def soft_shadow(shape, S):
    sh = S["shadow"]
    xml = (f'<a:effectLst xmlns:a="{A_NS}"><a:outerShdw blurRad="{sh["blur"]}" dist="{sh["dist"]}" '
           f'dir="{sh["dir"]}" algn="bl" rotWithShape="0">'
           f'<a:srgbClr val="000000"><a:alpha val="{sh["alpha"]}"/></a:srgbClr>'
           f'</a:outerShdw></a:effectLst>')
    spPr = _spPr(shape)
    for el in spPr.findall(_a("effectLst")):
        spPr.remove(el)
    spPr.append(etree.fromstring(xml))


def no_line(shape):
    shape.line.fill.background()


def set_adjustments_raw(shape, values):
    """프리셋 도형의 avLst를 원시 값으로 지정. values: {'adj1': 16200000, ...}"""
    spPr = _spPr(shape)
    prstGeom = spPr.find(_a("prstGeom"))
    avLst = prstGeom.find(_a("avLst"))
    if avLst is None:
        avLst = etree.SubElement(prstGeom, _a("avLst"))
    for el in list(avLst):
        avLst.remove(el)
    for name, v in values.items():
        gd = etree.SubElement(avLst, _a("gd"))
        gd.set("name", name)
        gd.set("fmla", f"val {int(v)}")


def _is_latin(text):
    return all(ord(ch) < 0x2100 for ch in text)


def add_text(slide, ctx, x, y, w, h, text, size, color, bold=False,
             align="l", anchor="t", spacing=None, wrap=True):
    """투명 텍스트박스. text 의 \n 은 문단 구분. cm 단위 좌표."""
    tb = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    aligns = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}
    lines = str(text).split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = aligns[align]
        if spacing:
            p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        f = r.font
        f.size = Pt(size)
        f.bold = bold
        f.color.rgb = color if isinstance(color, RGBColor) else RGBColor.from_string(color)
        fname = ctx["latin_font"] if _is_latin(line) and line.strip() else ctx["font"]
        f.name = fname
        rPr = r._r.get_or_add_rPr()
        for tag in ("ea", "cs"):
            el = rPr.find(_a(tag))
            if el is None:
                el = etree.SubElement(rPr, _a(tag))
            el.set("typeface", ctx["font"])
    return tb


def add_shape(slide, kind, x, y, w, h):
    return slide.shapes.add_shape(kind, Cm(x), Cm(y), Cm(w), Cm(h))


def card(slide, S, x, y, w, h, radius=None):
    """흰색 라운드 카드 + 얇은 테두리 + 소프트 섀도우 (레퍼런스 카드 스타일)"""
    sp = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sp.adjustments[0] = radius if radius is not None else S["corner_radius"]["card"]
    sp.fill.solid()
    sp.fill.fore_color.rgb = col(S, "white")
    sp.line.color.rgb = RGBColor.from_string(S["card_line"]["color"])
    sp.line.width = Emu(S["card_line"]["width_emu"])
    soft_shadow(sp, S)
    return sp


def gradient_card(slide, S, x, y, w, h, grad="diagonal", radius=None):
    sp = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    sp.adjustments[0] = radius if radius is not None else S["corner_radius"]["card"]
    g = S["gradient"][grad]
    gradient_fill(sp, g["colors"], g["angle"])
    no_line(sp)
    soft_shadow(sp, S)
    return sp


def ring(slide, S, cx, cy, d, grad="diagonal", thickness=None):
    """그라데이션 도넛 링. 슬라이드 밖으로 걸쳐도 된다 (음수 좌표 허용)."""
    sp = slide.shapes.add_shape(MSO_SHAPE.DONUT, Cm(cx - d / 2), Cm(cy - d / 2), Cm(d), Cm(d))
    sp.adjustments[0] = thickness if thickness is not None else S["motif"]["ring_thickness"]
    g = S["gradient"][grad]
    gradient_fill(sp, g["colors"], g["angle"])
    no_line(sp)
    return sp


def accent_circle(slide, S, cx, cy, d, grad="main"):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Cm(cx - d / 2), Cm(cy - d / 2), Cm(d), Cm(d))
    g = S["gradient"][grad]
    gradient_fill(sp, g["colors"], g["angle"])
    no_line(sp)
    return sp


def divider(slide, S, x, y, w):
    ln = slide.shapes.add_connector(1, Cm(x), Cm(y), Cm(x + w), Cm(y))
    ln.line.color.rgb = RGBColor.from_string(S["palette"]["divider"])
    ln.line.width = Emu(9525)
    return ln


def bottom_wave(slide, S, height=None, grad="main"):
    """하단 웨이브 — 사인 곡선 근사 프리폼"""
    h = height if height is not None else S["motif"]["wave_height_cm"]
    base = SH - h
    amp = h * 0.45
    n = 48
    x0, y0 = 0.0, base + amp * 0.5
    fb = slide.shapes.build_freeform(Emu(Cm(x0)), Emu(Cm(y0)), scale=1.0)
    pts = []
    for i in range(1, n + 1):
        x = SW * i / n
        y = base + amp * 0.5 * (1 - math.cos(2 * math.pi * x / SW)) - amp * 0.5 * math.sin(math.pi * x / SW)
        pts.append((Emu(Cm(x)), Emu(Cm(y))))
    pts.append((Emu(Cm(SW)), Emu(Cm(SH))))
    pts.append((Emu(Cm(0)), Emu(Cm(SH))))
    fb.add_line_segments(pts, close=True)
    sp = fb.convert_to_shape()
    g = S["gradient"][grad]
    gradient_fill(sp, g["colors"], g["angle"])
    no_line(sp)
    return sp


# ---------------------------------------------------------------- 이미지 처리

def _center_crop(pic, target_ratio):
    """세로/가로 비율에 맞게 중앙 크롭 (절대 늘리지 않음)"""
    iw, ih = pic.image.size
    native = iw / ih
    if native > target_ratio:
        crop = (1 - target_ratio / native) / 2
        pic.crop_left = crop
        pic.crop_right = crop
    elif native < target_ratio:
        crop = (1 - native / target_ratio) / 2
        pic.crop_top = crop
        pic.crop_bottom = crop


def _set_pic_geometry(pic, prst, adj=None):
    spPr = pic._element.spPr
    for el in spPr.findall(_a("prstGeom")) + spPr.findall(_a("custGeom")):
        spPr.remove(el)
    geom = etree.fromstring(f'<a:prstGeom xmlns:a="{A_NS}" prst="{prst}"><a:avLst/></a:prstGeom>')
    if adj is not None:
        gd = etree.SubElement(geom.find(_a("avLst")), _a("gd"))
        gd.set("name", "adj")
        gd.set("fmla", f"val {int(adj * 100000)}")
    xfrm = spPr.find(_a("xfrm"))
    if xfrm is not None:
        xfrm.addnext(geom)
    else:
        spPr.insert(0, geom)


def placeholder_graphic(slide, S, x, y, w, h, shape="rect", grad="diagonal"):
    """이미지가 없을 때: 그라데이션 도형 + 은은한 밝은 원 (텍스트 없음)"""
    if shape == "circle":
        sp = add_shape(slide, MSO_SHAPE.OVAL, x, y, w, h)
    else:
        sp = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        sp.adjustments[0] = S["corner_radius"]["card"]
    g = S["gradient"][grad]
    gradient_fill(sp, g["colors"], g["angle"])
    no_line(sp)
    d = min(w, h) * 0.42
    inner = add_shape(slide, MSO_SHAPE.OVAL, x + w * 0.58 - d / 2, y + h * 0.34 - d / 2, d, d)
    solid_fill_alpha(inner, S["palette"]["white"], 22000)
    no_line(inner)
    d2 = d * 0.45
    inner2 = add_shape(slide, MSO_SHAPE.OVAL, x + w * 0.30 - d2 / 2, y + h * 0.66 - d2 / 2, d2, d2)
    solid_fill_alpha(inner2, S["palette"]["white"], 16000)
    no_line(inner2)
    return sp


def image_slot(slide, S, path, x, y, w, h, shape="rect", rounded=True, grad="diagonal"):
    """이미지 경로가 유효하면 중앙 크롭해 배치, 없으면 플레이스홀더."""
    if path and Path(path).exists():
        pic = slide.shapes.add_picture(str(path), Cm(x), Cm(y), Cm(w), Cm(h))
        _center_crop(pic, w / h)
        if shape == "circle":
            _set_pic_geometry(pic, "ellipse")
        elif rounded:
            _set_pic_geometry(pic, "roundRect", adj=S["corner_radius"]["card"])
        return pic
    if path:
        print(f"[경고] 이미지를 찾을 수 없어 플레이스홀더로 대체: {path}")
    return placeholder_graphic(slide, S, x, y, w, h, shape=shape, grad=grad)


def progress_ring(slide, S, cx, cy, d, pct):
    """게이지 링: 회색 도넛 + 그라데이션 파이 + 흰 원 덮개"""
    t = 0.16
    base = slide.shapes.add_shape(MSO_SHAPE.DONUT, Cm(cx - d / 2), Cm(cy - d / 2), Cm(d), Cm(d))
    base.adjustments[0] = t
    base.fill.solid()
    base.fill.fore_color.rgb = RGBColor.from_string(S["palette"]["light"])
    no_line(base)
    pct = max(0.0, min(1.0, pct))
    if pct > 0.001:
        pie = slide.shapes.add_shape(MSO_SHAPE.PIE, Cm(cx - d / 2), Cm(cy - d / 2), Cm(d), Cm(d))
        start = 16200000  # 270도 = 12시 방향
        end = int((270 + 360 * pct) % 360 * 60000)
        set_adjustments_raw(pie, {"adj1": start, "adj2": end})
        g = S["gradient"]["diagonal"]
        gradient_fill(pie, g["colors"], g["angle"])
        no_line(pie)
    inner = d * (1 - 2 * t)
    cover = add_shape(slide, MSO_SHAPE.OVAL, cx - inner / 2, cy - inner / 2, inner, inner)
    cover.fill.solid()
    cover.fill.fore_color.rgb = col(S, "white")
    no_line(cover)


def add_footer(slide, S, ctx, idx):
    ts = S["type_scale"]
    add_text(slide, ctx, S["margins"]["left_cm"], SH - 0.95, 15, 0.6,
             ctx["deck_title"], ts["footer"], S["palette"]["footer"])
    add_text(slide, ctx, SW - 3.2, SH - 1.0, 1.6, 0.6, f"{idx:02d}",
             10, col(S, "primary"), bold=True, align="r")


def header(slide, S, ctx, c, title_key="title"):
    """킥커 + 슬라이드 타이틀 + 서브타이틀 공통 헤더. 다음 y 를 반환."""
    ts = S["type_scale"]
    m = S["margins"]["left_cm"]
    y = S["margins"]["top_cm"]
    if c.get("kicker"):
        add_text(slide, ctx, m, y, 26, 0.7, c["kicker"].upper(), ts["kicker"],
                 col(S, "primary"), bold=True)
        y += 0.85
    if c.get(title_key):
        add_text(slide, ctx, m, y, 28, 1.9, c[title_key], ts["title"],
                 col(S, "charcoal"), bold=True)
        y += 1.9
    if c.get("subtitle"):
        add_text(slide, ctx, m, y, 26, 0.8, c["subtitle"], 13.5, col(S, "gray"))
        y += 0.9
    return y + 0.5


def limit(items, n, layout):
    items = items or []
    if len(items) > n:
        print(f"[경고] {layout}: 항목 {len(items)}개 중 앞의 {n}개만 사용합니다.")
    return items[:n]


# ---------------------------------------------------------------- 레이아웃 (구도는 코드, 색·폰트는 스타일)

def lay_cover(slide, c, S, ctx):
    mot = S["motif"]
    d = mot["cover_ring_cm"]
    ring(slide, S, SW - 2.5, 4.2, d, grad="diagonal")
    c1, c2 = mot["accent_circles_cm"]
    accent_circle(slide, S, SW - 8.2, SH - 3.4, c1)
    accent_circle(slide, S, SW - 5.9, SH - 1.9, c2)
    accent_circle(slide, S, 2.7, 2.5, c1)
    ts = S["type_scale"]
    m = S["margins"]["left_cm"]
    if c.get("org"):
        add_text(slide, ctx, m + 0.9, 2.1, 12, 0.8, c["org"], 13, col(S, "charcoal"), bold=True)
    if c.get("kicker"):
        add_text(slide, ctx, m, 6.3, 21, 1.0, c["kicker"], 15, col(S, "primary"), bold=True)
    add_text(slide, ctx, m, 7.3, 22, 3.0, c.get("title", ""), ts["cover_title"],
             col(S, "charcoal"), bold=True)
    y = 10.6
    if c.get("tagline"):
        add_text(slide, ctx, m, y, 20, 1.5, c["tagline"], 24, col(S, "bright"), bold=True)
        y += 2.0
    if c.get("subtitle"):
        add_text(slide, ctx, m, y, 19, 1.2, c["subtitle"], 15, col(S, "gray"))
    if c.get("meta"):
        add_text(slide, ctx, m, SH - 2.4, 23, 1.0, c["meta"], 12, col(S, "charcoal"), bold=True)


def lay_toc(slide, c, S, ctx):
    d = 12.0
    image_slot(slide, S, c.get("image"), 2.4, 3.6, d, d, shape="circle")
    accent_circle(slide, S, 14.3, 3.9, 1.3)
    ts = S["type_scale"]
    x = 17.8
    add_text(slide, ctx, x, 2.4, 12, 0.7, c.get("kicker", "CONTENTS"), ts["kicker"],
             col(S, "primary"), bold=True)
    add_text(slide, ctx, x, 3.25, 12, 1.8, c.get("title", "목차"), 34, col(S, "charcoal"), bold=True)
    items = limit(c.get("items"), 6, "toc")
    y = 6.0
    step = 1.83
    for i, it in enumerate(items, 1):
        add_text(slide, ctx, x, y, 2.3, 1.4, f"{i:02d}", ts["big_number"], col(S, "primary"), bold=True)
        add_text(slide, ctx, x + 2.4, y + 0.35, 12.2, 1.0, str(it), 15.5, col(S, "charcoal"), bold=True)
        if i < len(items):
            divider(slide, S, x, y + step - 0.26, 13.9)
        y += step


def lay_intro(slide, c, S, ctx):
    panel = add_shape(slide, MSO_SHAPE.RECTANGLE, SW - 2.4, 0, 2.4, SH)
    g = S["gradient"]["vertical"]
    gradient_fill(panel, g["colors"], g["angle"])
    no_line(panel)
    m = S["margins"]["left_cm"]
    ts = S["type_scale"]
    y = 4.2
    if c.get("kicker"):
        add_text(slide, ctx, m, y, 12, 0.7, c["kicker"].upper(), ts["kicker"], col(S, "primary"), bold=True)
        y += 1.0
    add_text(slide, ctx, m, y, 12, 1.8, c.get("name", ""), ts["title"], col(S, "charcoal"), bold=True)
    y += 2.0
    if c.get("role"):
        add_text(slide, ctx, m, y, 12, 0.9, c["role"], 14, col(S, "bright"), bold=True)
        y += 1.3
    divider(slide, S, m, y, 10.5)
    y += 0.6
    if c.get("text"):
        add_text(slide, ctx, m, y, 11.5, 4.5, c["text"], 12.5, col(S, "gray"), spacing=1.35)
    cx, cy, cw, ch = 15.6, 3.0, 9.4, 13.2
    card(slide, S, cx, cy, cw, ch)
    image_slot(slide, S, c.get("image"), cx + 0.55, cy + 0.55, cw - 1.1, ch - 1.1)
    accent_circle(slide, S, cx + cw + 0.4, cy + 0.4, 1.2)


def lay_overview(slide, c, S, ctx):
    y0 = header(slide, S, ctx, c)
    m = S["margins"]["left_cm"]
    blocks = limit(c.get("blocks"), 3, "overview")
    y = max(y0, 5.6)
    for b in blocks:
        add_text(slide, ctx, m, y, 14.5, 0.9, b.get("head", ""), 15, col(S, "charcoal"), bold=True)
        add_text(slide, ctx, m, y + 0.9, 14.5, 1.6, b.get("desc", ""), 11.5, col(S, "gray"), spacing=1.25)
        y += 3.1
    cards = limit(c.get("cards"), 3, "overview") or ["", "", ""]
    cw, chh = 11.4, 3.6
    positions = [(19.4, 5.0), (20.6, 8.15), (19.4, 11.3)]
    for i, (cx, cy) in enumerate(positions[:len(cards)]):
        if i == 1:
            gradient_card(slide, S, cx, cy, cw, chh)
            txt_col = col(S, "white")
        else:
            card(slide, S, cx, cy, cw, chh)
            txt_col = col(S, "charcoal")
        if cards[i]:
            add_text(slide, ctx, cx + 0.9, cy, cw - 1.8, chh, str(cards[i]), 14,
                     txt_col, bold=True, anchor="m")


def lay_three_pillars(slide, c, S, ctx):
    header(slide, S, ctx, c)
    pillars = limit(c.get("pillars"), 3, "three_pillars")
    bottom_wave(slide, S)
    d = 6.4
    centers = [8.4, 16.93, 25.5][:len(pillars)]
    for cx, p in zip(centers, pillars):
        image_slot(slide, S, p.get("image"), cx - d / 2, 5.5, d, d, shape="circle")
        add_text(slide, ctx, cx - 4.2, 12.3, 8.4, 0.9, p.get("head", ""), 17,
                 col(S, "charcoal"), bold=True, align="c")
        add_text(slide, ctx, cx - 4.2, 13.3, 8.4, 1.8, p.get("desc", ""), 11.5,
                 col(S, "gray"), align="c", spacing=1.2)


def lay_team(slide, c, S, ctx):
    header(slide, S, ctx, c)
    band = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.8, 6.1, SW - 3.6, 10.4)
    band.adjustments[0] = 0.06
    g = S["gradient"]["soft"]
    gradient_fill(band, g["colors"], g["angle"])
    no_line(band)
    members = limit(c.get("members"), 4, "team")
    n = len(members) or 1
    d = 4.9
    span = (SW - 3.6) / n
    for i, mb in enumerate(members):
        cx = 1.8 + span * (i + 0.5)
        image_slot(slide, S, mb.get("image"), cx - d / 2, 7.3, d, d, shape="circle")
        add_text(slide, ctx, cx - span / 2 + 0.4, 12.6, span - 0.8, 0.8, mb.get("name", ""),
                 15, col(S, "charcoal"), bold=True, align="c")
        add_text(slide, ctx, cx - span / 2 + 0.4, 13.5, span - 0.8, 0.7, mb.get("role", ""),
                 11.5, col(S, "primary"), bold=True, align="c")


def lay_two_column(slide, c, S, ctx):
    header(slide, S, ctx, c)
    cols = limit(c.get("columns"), 2, "two_column")
    cw, chh, y = 14.7, 9.6, 5.7
    xs = [1.8, 17.35][:len(cols)]
    for x, cc in zip(xs, cols):
        card(slide, S, x, y, cw, chh)
        bar = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x + 0.7, y - 0.7, cw - 1.4, 1.9)
        bar.adjustments[0] = S["corner_radius"]["bar"]
        g = S["gradient"]["main"]
        gradient_fill(bar, g["colors"], g["angle"])
        no_line(bar)
        soft_shadow(bar, S)
        add_text(slide, ctx, x + 1.4, y - 0.7, cw - 2.8, 1.9, cc.get("head", ""), 16,
                 col(S, "white"), bold=True, anchor="m")
        dot = add_shape(slide, MSO_SHAPE.OVAL, x + 1.4, y + 2.4, 0.35, 0.35)
        dot.fill.solid()
        dot.fill.fore_color.rgb = col(S, "primary")
        no_line(dot)
        add_text(slide, ctx, x + 1.4, y + 3.2, cw - 2.8, 5.6, cc.get("desc", ""), 12.5,
                 col(S, "gray"), spacing=1.35)


def lay_process(slide, c, S, ctx):
    header(slide, S, ctx, c)
    steps = limit(c.get("steps"), 3, "process")
    ts = S["type_scale"]
    xs = [1.8, 12.5, 23.2][:len(steps)]
    for i, (x, st) in enumerate(zip(xs, steps)):
        mid = (i == 1 and len(steps) == 3)
        y, hh = (5.5, 10.2) if mid else (6.1, 9.0)
        if mid:
            gradient_card(slide, S, x, y, 8.9, hh)
            num_c, head_c, body_c = col(S, "white"), col(S, "white"), RGBColor.from_string("FFF3E0")
        else:
            card(slide, S, x, y, 8.9, hh)
            num_c, head_c, body_c = col(S, "primary"), col(S, "charcoal"), col(S, "gray")
        add_text(slide, ctx, x + 1.0, y + 0.9, 6.9, 1.8, f"{i + 1:02d}", 38, num_c, bold=True)
        add_text(slide, ctx, x + 1.0, y + 3.0, 6.9, 1.0, st.get("head", ""), ts["card_head"],
                 head_c, bold=True)
        add_text(slide, ctx, x + 1.0, y + 4.2, 6.9, hh - 5.0, st.get("desc", ""), 12,
                 body_c, spacing=1.3)


def lay_four_cards(slide, c, S, ctx):
    header(slide, S, ctx, c)
    cards_ = limit(c.get("cards"), 4, "four_cards")
    cw, y, hh = 7.1, 5.8, 9.2
    gap = (SW - 3.6 - cw * 4) / 3 if len(cards_) == 4 else 0.7
    for i, cd in enumerate(cards_):
        x = 1.8 + i * (cw + gap)
        card(slide, S, x, y, cw, hh)
        top = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, cw, 2.5)
        top.adjustments[0] = 0.11
        g = S["gradient"]["main"]
        gradient_fill(top, g["colors"], g["angle"])
        no_line(top)
        add_text(slide, ctx, x + 0.6, y, cw - 1.2, 2.5, cd.get("head", ""), 14,
                 col(S, "white"), bold=True, anchor="m")
        add_text(slide, ctx, x + 0.7, y + 3.2, cw - 1.4, hh - 4.0, cd.get("desc", ""), 11.5,
                 col(S, "gray"), spacing=1.3)


def lay_vision(slide, c, S, ctx):
    ring(slide, S, SW + 0.6, 13.5, 9.5, grad="diagonal")
    y0 = header(slide, S, ctx, c)
    image_slot(slide, S, c.get("image"), 1.8, max(y0, 5.2), 13.2, 11.0)
    if c.get("badge"):
        bd = add_shape(slide, MSO_SHAPE.OVAL, 13.4, 12.4, 4.4, 4.4)
        bd.fill.solid()
        bd.fill.fore_color.rgb = col(S, "white")
        no_line(bd)
        soft_shadow(bd, S)
        add_text(slide, ctx, 13.7, 12.4, 3.8, 4.4, c["badge"], 13, col(S, "primary"),
                 bold=True, align="c", anchor="m")
    if c.get("text"):
        add_text(slide, ctx, 18.9, max(y0, 5.2) + 1.2, 11.3, 8.5, c["text"], 15,
                 col(S, "gray"), spacing=1.45)


def lay_gallery(slide, c, S, ctx):
    header(slide, S, ctx, c)
    imgs = limit(c.get("images"), 3, "gallery")
    imgs += [None] * (3 - len(imgs))
    image_slot(slide, S, imgs[0], 1.8, 5.3, 15.2, 10.9)
    image_slot(slide, S, imgs[1], 17.7, 5.3, 6.9, 5.2, grad="vertical")
    image_slot(slide, S, imgs[2], 25.3, 5.3, 6.9, 5.2, grad="main")
    cap = gradient_card(slide, S, 17.7, 11.0, 14.5, 5.2, grad="main")
    if c.get("caption"):
        add_text(slide, ctx, 18.8, 11.0, 12.3, 5.2, c["caption"], 14, col(S, "white"),
                 bold=True, anchor="m", spacing=1.3)


def lay_kpi(slide, c, S, ctx):
    header(slide, S, ctx, c)
    an = c.get("analysis") or {}
    gradient_card(slide, S, 1.8, 5.4, 13.0, 10.6)
    add_text(slide, ctx, 3.0, 6.6, 10.6, 1.2, an.get("head", ""), 18, col(S, "white"), bold=True)
    y = 8.6
    for b in limit(an.get("bullets"), 3, "kpi"):
        dot = add_shape(slide, MSO_SHAPE.OVAL, 3.0, y + 0.22, 0.3, 0.3)
        dot.fill.solid()
        dot.fill.fore_color.rgb = col(S, "white")
        no_line(dot)
        add_text(slide, ctx, 3.7, y, 10.0, 1.8, str(b), 12, RGBColor.from_string("FFF3E0"),
                 spacing=1.2)
        y += 2.2
    rings = limit(c.get("rings"), 3, "kpi")
    n = len(rings) or 1
    span = (SW - 16.6) / n
    d = 6.2
    for i, rg in enumerate(rings):
        cx = 16.6 + span * (i + 0.5)
        cy = 9.0
        progress_ring(slide, S, cx, cy, d, float(rg.get("value", 0)) / 100.0)
        add_text(slide, ctx, cx - 2.4, cy - 0.75, 4.8, 1.5, f"{int(rg.get('value', 0))}%",
                 20, col(S, "primary"), bold=True, align="c", anchor="m")
        add_text(slide, ctx, cx - span / 2 + 0.3, cy + d / 2 + 0.6, span - 0.6, 1.5,
                 rg.get("label", ""), 12.5, col(S, "charcoal"), bold=True, align="c")


def lay_report(slide, c, S, ctx):
    header(slide, S, ctx, c)
    ix, iy, iw, ih = 1.8, 5.3, 13.6, 10.7
    image_slot(slide, S, c.get("image"), ix, iy, iw, ih)
    for (bx, by, bw, bh) in [(ix - 0.35, iy - 0.35, 3.4, 0.35), (ix - 0.35, iy - 0.35, 0.35, 3.4),
                             (ix + iw - 3.05, iy + ih, 3.4, 0.35), (ix + iw, iy + ih - 3.05, 0.35, 3.4)]:
        bar = add_shape(slide, MSO_SHAPE.RECTANGLE, bx, by, bw, bh)
        g = S["gradient"]["main"]
        gradient_fill(bar, g["colors"], g["angle"])
        no_line(bar)
    x = 17.4
    y = 5.4
    findings = limit(c.get("findings"), 4, "report")
    step = 10.6 / max(len(findings), 1)
    for i, f in enumerate(findings, 1):
        add_text(slide, ctx, x, y, 2.2, 1.3, f"{i:02d}", 22, col(S, "primary"), bold=True)
        add_text(slide, ctx, x + 2.5, y + 0.15, 12.0, step - 0.6, str(f), 12.5,
                 col(S, "charcoal"), spacing=1.25)
        if i < len(findings):
            divider(slide, S, x, y + step - 0.35, 14.2)
        y += step


def lay_bar_chart(slide, c, S, ctx):
    header(slide, S, ctx, c)
    bars = limit(c.get("bars"), 6, "bar_chart")
    vals = [float(b.get("value", 0)) for b in bars]
    vmax = max(vals) if vals else 1
    x0, x1, base_y, max_h = 1.8, 20.0, 15.3, 7.6
    n = len(bars) or 1
    slot = (x1 - x0) / n
    bw = min(slot * 0.52, 2.6)
    for i, (b, v) in enumerate(zip(bars, vals)):
        bh = max_h * (v / vmax) if vmax else 0
        bx = x0 + slot * i + (slot - bw) / 2
        sp = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, bx, base_y - bh, bw, max(bh, 0.3))
        sp.adjustments[0] = 0.25
        is_max = (v == vmax)
        g = S["gradient"]["vertical"] if not is_max else {"colors": [S["palette"]["primary"], S["palette"]["bright"]], "angle": 90}
        gradient_fill(sp, g["colors"], g["angle"])
        no_line(sp)
        add_text(slide, ctx, bx - 1.2, base_y - bh - 0.85, bw + 2.4, 0.7, f"{v:g}",
                 12 if is_max else 11, col(S, "primary") if is_max else col(S, "gray"),
                 bold=True, align="c")
        add_text(slide, ctx, bx - 1.5, base_y + 0.3, bw + 3.0, 1.2, b.get("label", ""),
                 10.5, col(S, "charcoal"), align="c")
    divider(slide, S, x0, base_y, x1 - x0)
    sm = c.get("summary") or {}
    card(slide, S, 21.4, 6.4, 10.7, 6.6)
    add_text(slide, ctx, 22.5, 7.4, 8.5, 2.0, sm.get("stat", ""), 32, col(S, "primary"), bold=True)
    add_text(slide, ctx, 22.5, 9.6, 8.5, 3.0, sm.get("text", ""), 12, col(S, "gray"), spacing=1.3)


def lay_full_image(slide, c, S, ctx):
    img = c.get("image")
    if img and Path(img).exists():
        pic = slide.shapes.add_picture(str(img), 0, 0, Emu(SLIDE_W_EMU), Emu(SLIDE_H_EMU))
        _center_crop(pic, SW / SH)
        ov = add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
        solid_fill_alpha(ov, "000000", 45000)
        no_line(ov)
    else:
        bg = add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
        g = S["gradient"]["dark"]
        gradient_fill(bg, g["colors"], g["angle"])
        no_line(bg)
        ring(slide, S, SW - 1.5, 2.0, 10.0)
    pw, phh = 21.0, 5.4
    panel = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, (SW - pw) / 2, 11.4, pw, phh)
    panel.adjustments[0] = 0.09
    g = S["gradient"]["main"]
    gradient_fill(panel, g["colors"], g["angle"], alphas=[82000, 82000])
    no_line(panel)
    add_text(slide, ctx, (SW - pw) / 2 + 1.5, 12.1, pw - 3.0, 1.6, c.get("title", ""), 24,
             col(S, "white"), bold=True, align="c")
    add_text(slide, ctx, (SW - pw) / 2 + 1.5, 13.9, pw - 3.0, 2.4, c.get("text", ""), 13,
             col(S, "white"), align="c", spacing=1.3)


def lay_testimonial(slide, c, S, ctx):
    header(slide, S, ctx, c)
    band = gradient_card(slide, S, 1.8, 7.6, SW - 3.6, 7.4, grad="main", radius=0.1)
    quotes = limit(c.get("quotes"), 2, "testimonial")
    half = (SW - 3.6) / 2
    d = 4.4
    for i, q in enumerate(quotes):
        x = 1.8 + half * i
        cx = x + 3.6
        image_slot(slide, S, q.get("image"), cx - d / 2, 5.9, d, d, shape="circle", grad="soft")
        add_text(slide, ctx, x + 6.3, 8.6, half - 7.4, 3.4, f"“{q.get('text', '')}”",
                 12.5, col(S, "white"), spacing=1.3)
        add_text(slide, ctx, x + 6.3, 12.4, half - 7.4, 0.8,
                 f"{q.get('name', '')}  ·  {q.get('role', '')}", 11.5,
                 RGBColor.from_string("FFF0D0"), bold=True)


def lay_case_study(slide, c, S, ctx):
    header(slide, S, ctx, c)
    cases = limit(c.get("cases"), 3, "case_study")
    n = len(cases) or 1
    gap = 0.9
    cw = (SW - 3.6 - gap * (n - 1)) / n
    grads = ["diagonal", "vertical", "main"]
    for i, cs in enumerate(cases):
        x = 1.8 + i * (cw + gap)
        image_slot(slide, S, cs.get("image"), x, 5.5, cw, 8.6, grad=grads[i % 3])
        dot = add_shape(slide, MSO_SHAPE.OVAL, x + 0.1, 14.75, 0.35, 0.35)
        dot.fill.solid()
        dot.fill.fore_color.rgb = col(S, "primary")
        no_line(dot)
        add_text(slide, ctx, x + 0.7, 14.55, cw - 0.8, 1.6, cs.get("caption", ""), 12.5,
                 col(S, "charcoal"), bold=True, spacing=1.2)


def lay_section(slide, c, S, ctx):
    ring(slide, S, SW + 1.0, SH - 1.0, 12.0, grad="diagonal")
    accent_circle(slide, S, SW - 6.0, 3.0, 1.2)
    m = S["margins"]["left_cm"]
    if c.get("number"):
        add_text(slide, ctx, m, 4.6, 12, 4.5, str(c["number"]), 88, col(S, "primary"), bold=True)
    bar = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, m + 0.1, 10.4, 3.4, 0.5)
    bar.adjustments[0] = 0.5
    g = S["gradient"]["main"]
    gradient_fill(bar, g["colors"], g["angle"])
    no_line(bar)
    add_text(slide, ctx, m, 11.4, 26, 2.2, c.get("title", ""), 36, col(S, "charcoal"), bold=True)
    if c.get("subtitle"):
        add_text(slide, ctx, m, 13.7, 24, 1.2, c["subtitle"], 14, col(S, "gray"))


def lay_thank_you(slide, c, S, ctx):
    mot = S["motif"]
    ring(slide, S, SW - 2.0, SH - 4.5, mot["cover_ring_cm"] * 0.9, grad="diagonal")
    c1, c2 = mot["accent_circles_cm"]
    accent_circle(slide, S, SW - 9.5, 3.2, c1)
    accent_circle(slide, S, SW - 7.6, 5.0, c2)
    m = S["margins"]["left_cm"]
    add_text(slide, ctx, m, 6.6, 24, 3.2, c.get("title", "THANK YOU"), S["type_scale"]["cover_title"],
             col(S, "charcoal"), bold=True)
    if c.get("subtitle"):
        add_text(slide, ctx, m, 10.4, 20, 1.6, c["subtitle"], 15, col(S, "gray"), spacing=1.3)
    if c.get("contact"):
        divider(slide, S, m, 14.4, 11.0)
        add_text(slide, ctx, m, 14.9, 20, 1.2, c["contact"], 12, col(S, "charcoal"), bold=True)


LAYOUTS = {
    "cover": lay_cover, "toc": lay_toc, "intro": lay_intro, "overview": lay_overview,
    "three_pillars": lay_three_pillars, "team": lay_team, "two_column": lay_two_column,
    "process": lay_process, "four_cards": lay_four_cards, "vision": lay_vision,
    "gallery": lay_gallery, "kpi": lay_kpi, "report": lay_report, "bar_chart": lay_bar_chart,
    "full_image": lay_full_image, "testimonial": lay_testimonial, "case_study": lay_case_study,
    "section": lay_section, "thank_you": lay_thank_you,
}

NO_FOOTER = {"cover", "section", "full_image", "thank_you"}


# ---------------------------------------------------------------- 빌드

def build(spec_path, out_path, style_name=None):
    with open(spec_path, encoding="utf-8") as f:
        spec = json.load(f)
    style_name = style_name or spec.get("style") or "orange-gradient"
    S = load_style(style_name)

    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W_EMU)
    prs.slide_height = Emu(SLIDE_H_EMU)
    blank = prs.slide_layouts[6]

    ctx = {
        "deck_title": spec.get("title", ""),
        "font": spec.get("font") or S["font"]["main"],
        "latin_font": S["font"]["latin"],
    }

    for idx, sl in enumerate(spec.get("slides", []), 1):
        name = sl.get("layout", "")
        if name not in LAYOUTS:
            valid = ", ".join(sorted(LAYOUTS))
            sys.exit(f"[오류] 슬라이드 {idx}: 알 수 없는 레이아웃 '{name}'.\n사용 가능한 레이아웃: {valid}")
        slide = prs.slides.add_slide(blank)
        LAYOUTS[name](slide, sl, S, ctx)
        if name not in NO_FOOTER:
            add_footer(slide, S, ctx, idx)
        if sl.get("notes"):
            slide.notes_slide.notes_text_frame.text = sl["notes"]

    prs.save(out_path)
    print(f"완료: {out_path} ({len(spec.get('slides', []))}장, 스타일: {style_name})")


def main():
    ap = argparse.ArgumentParser(description="TextbookBuilder PPT 생성기")
    ap.add_argument("spec", help="슬라이드 spec JSON 경로")
    ap.add_argument("-o", "--output", required=True, help="출력 .pptx 경로")
    ap.add_argument("--style", default=None, help="ppt-styles/<이름>.json 스타일 (기본: orange-gradient)")
    args = ap.parse_args()
    build(args.spec, args.output, args.style)


if __name__ == "__main__":
    main()
